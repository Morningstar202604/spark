"""办公文档工具：读写 docx / xlsx / pdf，补全「办公 Agent」能力（对标钉钉/通义/WPS AI）。

依赖：python-docx / openpyxl / pypdf（纯 Python，可选；缺失时工具返回明确提示，不影响其它工具）。
读取类归 read（免审批），写入类归 write（suggest 模式询问）。
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

from spark.tools.base import Tool, ToolContext, resolve_path, is_within

_HAS_DOCX = importlib.util.find_spec("docx") is not None
_HAS_XLSX = importlib.util.find_spec("openpyxl") is not None
_HAS_PDF = importlib.util.find_spec("pypdf") is not None


def _need(lib: str) -> str:
    return f"错误：未安装 {lib} 库，无法执行此工具。请先运行：pip install {lib}"


async def _read_docx(args: dict, ctx: ToolContext) -> str:
    if not _HAS_DOCX:
        return _need("python-docx")
    path = resolve_path(str(args.get("path", "")), ctx.workdir)
    if not path.exists():
        return f"错误：文件不存在 {path}"
    try:
        from docx import Document  # noqa: PLC0415

        doc = Document(path)
        parts: list[str] = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                parts.append(t)
        for tbl in doc.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(cells))
        text = "\n".join(parts)
        if not text.strip():
            return f"（{path} 无可提取文本）"
        return f"# {path}\n\n{text}"
    except Exception as e:  # noqa: BLE001
        return f"读取 docx 失败：{e}"


async def _write_docx(args: dict, ctx: ToolContext) -> str:
    if not _HAS_DOCX:
        return _need("python-docx")
    path = resolve_path(str(args.get("path", "")), ctx.workdir)
    if not is_within(path, ctx.workdir):
        return "错误：写入路径需在工作目录内"
    content = str(args.get("content", ""))
    try:
        from docx import Document  # noqa: PLC0415
        from docx.shared import Pt  # noqa: PLC0415

        doc = Document()
        for line in content.split("\n"):
            line = line.rstrip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        return f"已生成 Word 文档：{path}（{content.count(chr(10)) + 1} 行）"
    except Exception as e:  # noqa: BLE001
        return f"生成 docx 失败：{e}"


async def _read_xlsx(args: dict, ctx: ToolContext) -> str:
    if not _HAS_XLSX:
        return _need("openpyxl")
    path = resolve_path(str(args.get("path", "")), ctx.workdir)
    if not path.exists():
        return f"错误：文件不存在 {path}"
    try:
        from openpyxl import load_workbook  # noqa: PLC0415

        wb = load_workbook(path, data_only=True)
        out: list[str] = []
        for ws in wb.worksheets:
            out.append(f"## 工作表：{ws.title}（{ws.max_row} 行 × {ws.max_column} 列）")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    out.append(" | ".join(cells))
        return "\n".join(out)
    except Exception as e:  # noqa: BLE001
        return f"读取 xlsx 失败：{e}"


async def _write_xlsx(args: dict, ctx: ToolContext) -> str:
    if not _HAS_XLSX:
        return _need("openpyxl")
    path = resolve_path(str(args.get("path", "")), ctx.workdir)
    if not is_within(path, ctx.workdir):
        return "错误：写入路径需在工作目录内"
    headers = args.get("headers") or []
    rows = args.get("rows") or []
    if not isinstance(rows, list) or not rows:
        return "错误：rows 需为非空二维数组"
    sheet = str(args.get("sheet") or "Sheet1")
    try:
        from openpyxl import Workbook  # noqa: PLC0415

        wb = Workbook()
        ws = wb.active
        ws.title = sheet[:31] or "Sheet1"
        if headers:
            ws.append([str(h) for h in headers])
        for r in rows:
            ws.append([str(c) if c is not None else "" for c in r])
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return f"已生成 Excel：{path}（{len(rows)} 行数据）"
    except Exception as e:  # noqa: BLE001
        return f"生成 xlsx 失败：{e}"


async def _read_pdf(args: dict, ctx: ToolContext) -> str:
    if not _HAS_PDF:
        return _need("pypdf")
    path = resolve_path(str(args.get("path", "")), ctx.workdir)
    if not path.exists():
        return f"错误：文件不存在 {path}"
    page_lim = max(1, min(int(args.get("max_pages", 20)), 100))
    try:
        from pypdf import PdfReader  # noqa: PLC0415

        reader = PdfReader(path)
        total = len(reader.pages)
        parts: list[str] = []
        for i, page in enumerate(reader.pages):
            if i >= page_lim:
                parts.append(f"…（共 {total} 页，仅展示前 {page_lim} 页）")
                break
            parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
        if not text:
            return f"（{path} 无可提取文本，可能是扫描件，需要 OCR）"
        return f"# {path}（{total} 页，已读 {min(page_lim, total)} 页）\n\n{text}"
    except Exception as e:  # noqa: BLE001
        return f"读取 pdf 失败：{e}"


def build_office_tools() -> list[Tool]:
    return [
        Tool(
            name="read_docx",
            description="读取 Word(.docx) 文档的段落和表格文本，返回纯文本。用于查看/摘要/引用 Word 文件。",
            parameters={
                "type": "object",
                "properties": {"path": {"type": "string", "description": ".docx 文件路径（相对或绝对）"}},
                "required": ["path"],
            },
            category="read",
            handler=_read_docx,
        ),
        Tool(
            name="write_docx",
            description="生成 Word(.docx) 文档。content 支持简单标记：# 一级标题 / ## 二级标题 / ### 三级标题 / - 列表项 / 普通段落。用于生成报告、方案、纪要等。",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "输出 .docx 文件路径"},
                    "content": {"type": "string", "description": "文档内容（# 标题 / ## 小节 / - 列表）"},
                },
                "required": ["path", "content"],
            },
            category="write",
            handler=_write_docx,
        ),
        Tool(
            name="read_xlsx",
            description="读取 Excel(.xlsx) 所有工作表，按表格形式返回（每个单元格 | 分隔）。用于查看/分析表格数据。",
            parameters={
                "type": "object",
                "properties": {"path": {"type": "string", "description": ".xlsx 文件路径"}},
                "required": ["path"],
            },
            category="read",
            handler=_read_xlsx,
        ),
        Tool(
            name="write_xlsx",
            description="生成 Excel(.xlsx)。rows 为二维数组（每行一个数组），headers 可选表头。用于产出表格/报表数据。",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "输出 .xlsx 文件路径"},
                    "rows": {"type": "array", "description": "数据行，[[列1,列2,...], ...]"},
                    "headers": {"type": "array", "description": "可选表头，如 [\"月份\",\"金额\"]"},
                    "sheet": {"type": "string", "description": "可选工作表名，默认 Sheet1"},
                },
                "required": ["path", "rows"],
            },
            category="write",
            handler=_write_xlsx,
        ),
        Tool(
            name="read_pdf",
            description="读取 PDF 文本内容（提取文字；扫描件/图片型 PDF 需 OCR 才能读取）。max_pages 控制读取页数。",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": ".pdf 文件路径"},
                    "max_pages": {"type": "integer", "description": "最多读取页数，默认 20，最大 100"},
                },
                "required": ["path"],
            },
            category="read",
            handler=_read_pdf,
        ),
    ]